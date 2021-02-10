import urllib.request, urllib.parse, urllib.error
from bs4 import BeautifulSoup
import pdfplumber
import pandas as pd
from fpdf import FPDF
from docx import Document
from docx.shared import Pt
import ssl
import re
import os
import codecs
import json
import xlsxwriter

emails_found = []

def file_extension(location):
	if location.lower().endswith(('.pdf')) == True:
		return "pdf"
	elif location.lower().endswith(('.txt')) == True:
		return "txt"
	elif location.lower().endswith(('.json')) == True:
		return "json"
	elif location.lower().endswith(('.txt', '.pdf', 'json')) == False:
		return False

def where_to_go(location, extension, index=0):
	if extension == "txt" and index==0:
		from_text(location)
	elif extension == "pdf" and index==0:
		from_pdf_or_docx(location)
	elif extension == "json" and index==0:
		from_json(location)
	elif index == 1:
		return emails_found

def internet_conection(url):
    try:
        urllib.request.urlopen(url, timeout=10)
        from_page(url)
    except:
    	return False

def from_pdf_or_docx(location):
	pdf = pdfplumber.open(location)
	f = open(f"{os.getcwd()[:2]}temporary.txt","a", encoding='utf-8')
	for page in range(len(pdf.pages)):
		text = pdf.pages[page].extract_text()
		try:
			cleaned_string = ''.join(c for c in text if valid_xml_char_ordinal(c))
			f.write(str(cleaned_string))
		except:
			f.write(str(text))
	pdf.close()
	f.close()
	from_text(f"{os.getcwd()[:2]}temporary.txt", 1)

def from_json(location):
	with open(location, encoding='utf-8-sig') as json_file:
		data = json.load(json_file)
	f = open(f"{os.getcwd()[:2]}temporary.txt","a", encoding='utf-8')
	for text in data:
		f.write(str(f"\n{text}"))
	f.close()
	from_text(f"{os.getcwd()[:2]}temporary.txt", 1)

def from_page(url):
	# Ignore SSL certificate errors
	ctx = ssl.create_default_context()
	ctx.check_hostname = False
	ctx.verify_mode = ssl.CERT_NONE

	html = urllib.request.urlopen(url, context=ctx).read()
	soup = BeautifulSoup(html, "html.parser")
	f = open(f"{os.getcwd()[:2]}temporary.txt","w+", encoding='utf-8')
	f.write(soup.prettify())
	f.close()
	from_text(f"{os.getcwd()[:2]}temporary.txt", 1)

def from_text(location, index=0):
	global emails_found
	if len(emails_found) > 0:
		emails_found = []
	file = codecs.open(location, encoding='utf-8-sig')
	for line in file:
		line = line.strip()
		if re.search('\S+@\S+[\w-]+\.+[\w-]{2,100}$', line):
			filtered = re.findall('\S+@\S+[\w-]+\.+[\w-]{2,100}$', line)
			for email in filtered:
				email = ''.join(email)
				emails_found.append(''.join(email))
	file.close()
	if index == 1:
		os.remove(location)

def emails_existence(emails_found):
	if len(emails_found) > 0:
		category(emails_found)
	else:
		return False

def category(emails_found, index=0):
	email_category = {"All": emails_found}
	for extension in emails_found:
		formats = re.findall('@\S+[\w-]+\.+[\w-]{2,100}$', extension)
		for Format in formats:
			email_category[Format] = []
	for email in emails_found:
		formats = re.findall('@\S+[\w-]+\.+[\w-]{2,100}$', email)
		for mail in formats:
			if mail in email:
				email_category[mail].append(email)
	if index == 1:
		return len(category(emails_found).keys())
	return email_category

def amount_of_emails_found(lenth_email_category):
	if (lenth_email_category-1) == 1:
		text = "Only one email domain was found"
		return text
	else:
		text = f"""{lenth_email_category-1} differents emails domains were found
Which one would you like to export?"""
		return text

def output_file(filtered_email_list, index, location, extension):
	global ending
	global name
	ending = extension
	for n in range(10000):
		if not (f"output_file_{n}{extension}") in os.listdir(location):
			name = (f"output_file_{n}")
			break
	full_path_location = f"{location}{location[2]}{name}{extension}"
	if index == 3:
		n = 0
		f = open(full_path_location,"w", encoding='utf-8-sig')
		for email in filtered_email_list:
			n+=1
			if n == 1:
				f.write(email)
			else:
				f.write(f"\n{email}")
		f.close()

	elif index == 4:
		df = pd.DataFrame({'':filtered_email_list}) 
		df.to_csv(full_path_location, index=False, encoding='utf-8-sig')

	elif index == 5:
		with open(full_path_location, 'w', encoding='utf-8-sig') as outfile:
			json.dump(filtered_email_list, outfile)

	elif index == 6:
		pdf = FPDF()
		pdf.add_page()
		pdf.set_font("Arial", size = 12)
		for email in filtered_email_list:
			pdf.cell(300, 6, txt = email, ln = 1, align = 'L')
		pdf.output(full_path_location)

	elif index == 7:
		book = xlsxwriter.Workbook(full_path_location)     
		sheet = book.add_worksheet() 
		row = 0    
		column = 0
		for email in filtered_email_list:
			sheet.write(row, column, email)
			row += 1
		book.close()

	elif index == 8:
		document = Document()
		for email in filtered_email_list:
			cleaned_string = ''.join(c for c in email if valid_xml_char_ordinal(c))
			paragraph = document.add_paragraph(cleaned_string)
			paragraph.paragraph_format.space_before = Pt(2)
			paragraph.paragraph_format.space_after = Pt(2)
		document.add_page_break()
		document.save(full_path_location)

def output_file_location(path, index=0):
	if index == 1:
		return path
	else:
		return path+path[2]+name+ending


def valid_xml_char_ordinal(c):
    codepoint = ord(c)
    return (
        0x20 <= codepoint <= 0xD7FF or
        codepoint in (0x9, 0xA, 0xD) or
        0xE000 <= codepoint <= 0xFFFD or
        0x10000 <= codepoint <= 0x10FFFF
        )
